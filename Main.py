import base64
import html
import io
import os
import re
import ssl
import time
from pathlib import Path
from typing import Any

import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from requests.adapters import HTTPAdapter
from urllib3 import PoolManager


DUPR_API = "https://api.dupr.gg/player/v1.0/search/public/"
DUPR_ID_API = "https://api.dupr.gg/player/search/byDuprId/"
DUPR_PLAYER_API = "https://api.dupr.gg/player/v1.0/{player_id}/"
ADVANCED_CUTOFF = 4.0
AGE_DIVISIONS = ["U18", "18-30", "Above 30"]
LEVELS = ["Advanced", "Intermediate", "No DUPR Rating"]
GENDERS = ["Boys / Men", "Girls / Women"]
APP_DIRECTORY = Path(__file__).resolve().parent

st.set_page_config(
    page_title="Pickleball Player Manager",
    page_icon="🏓",
    layout="wide",
    initial_sidebar_state="collapsed",
)


def apply_theme() -> None:
    st.markdown(
        """
        <style>
        :root {
            --ink: #102a43;
            --muted: #627d98;
            --navy: #12355b;
            --court: #118a74;
            --lime: #d7f171;
            --surface: #ffffff;
            --line: #dfe8ef;
        }
        .stApp {
            background: linear-gradient(180deg, #f4faf8 0, #f8fafc 300px, #f7f9fc 100%);
            color: var(--ink);
        }
        .block-container {padding-top: 1.2rem; padding-bottom: 3rem; max-width: 1440px;}
        .hero {
            background: linear-gradient(118deg, #0d6b61, #12355b 72%);
            color: white; border-radius: 22px; padding: 1.55rem 1.8rem;
            box-shadow: 0 14px 35px rgba(18, 53, 91, .16); margin-bottom: 1rem;
            position: relative; overflow: hidden;
        }
        .hero:after {content:""; position:absolute; width:180px; height:180px; border:30px solid rgba(215,241,113,.13); border-radius:50%; right:-48px; top:-80px;}
        .hero-kicker {font-size:.76rem; font-weight:800; letter-spacing:.15em; text-transform:uppercase; color:var(--lime); margin-bottom:.35rem;}
        .hero h1 {font-size:clamp(1.55rem, 4vw, 2.45rem); line-height:1.08; margin:0; color:white;}
        .hero p {font-size:clamp(.9rem, 2vw, 1.05rem); color:#dcefeb; margin:.55rem 0 0;}
        .hero h1, .hero p {padding-right:185px; position:relative; z-index:1;}
        .hero-kicker {position:relative; z-index:1;}
        .hero-logo-shell {
            position:absolute; z-index:2; right:1.65rem; top:50%; transform:translateY(-50%);
            width:150px; height:122px; padding:5px; border-radius:17px;
            background:rgba(255,255,255,.94); border:1px solid rgba(255,255,255,.6);
            box-shadow:0 10px 25px rgba(0,0,0,.2);
        }
        .hero-logo {width:100%; height:100%; display:block; object-fit:cover; border-radius:12px;}
        div[data-testid="stMetric"] {
            background: rgba(255,255,255,.94); border:1px solid var(--line); border-radius:16px;
            padding:1rem 1.05rem; box-shadow:0 5px 18px rgba(16,42,67,.055); min-height:108px;
        }
        div[data-testid="stMetricLabel"] {color:var(--muted); font-weight:700;}
        div[data-testid="stMetricValue"] {color:var(--navy); font-weight:800;}
        div[data-testid="stTabs"] button {font-weight:750; padding-left:.8rem; padding-right:.8rem;}
        div[data-testid="stTabs"] div[data-baseweb="tab-list"] {
            overflow-x:auto; overflow-y:hidden; scrollbar-width:thin; scroll-snap-type:x proximity;
        }
        div[data-testid="stTabs"] button {white-space:nowrap; scroll-snap-align:start; flex-shrink:0;}
        .section-card {background:white; border:1px solid var(--line); border-radius:18px; padding:1.05rem 1.2rem; margin:.35rem 0 1rem; box-shadow:0 5px 18px rgba(16,42,67,.045);}
        .eyebrow {color:var(--court); font-size:.74rem; font-weight:850; letter-spacing:.12em; text-transform:uppercase;}
        .result-card {background:linear-gradient(135deg,#fff,#f0faf7); border:1px solid #cde8e1; border-radius:20px; padding:1.25rem 1.35rem; box-shadow:0 8px 24px rgba(17,138,116,.09);}
        .result-name {font-size:1.45rem; font-weight:850; color:var(--navy); margin:.25rem 0;}
        .badge {display:inline-block; border-radius:999px; padding:.35rem .72rem; font-size:.76rem; font-weight:900; letter-spacing:.08em; margin:.25rem 0 .85rem;}
        .advanced {background:#dff7eb; color:#08724f; border:1px solid #94dfbf;}
        .intermediate {background:#e7f0ff; color:#2457a6; border:1px solid #aec9f3;}
        .manual {background:#fff1d6; color:#8a5500; border:1px solid #f2cc84;}
        .rating-grid {display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:.65rem; margin-top:.35rem;}
        .rating-box {background:white; border:1px solid var(--line); border-radius:13px; padding:.75rem;}
        .rating-label {color:var(--muted); font-size:.75rem; font-weight:700;}
        .rating-value {color:var(--navy); font-size:1.15rem; font-weight:850;}
        .status-pill {display:inline-block; padding:.22rem .55rem; background:#eef6f4; border-radius:999px; color:#176c60; font-size:.78rem; font-weight:700; margin:.12rem;}
        .stButton > button, .stDownloadButton > button {min-height:2.8rem; border-radius:11px; font-weight:750;}
        div[data-testid="stFileUploader"] {background:white; border-radius:16px; padding:.35rem;}
        div[data-testid="stDataFrame"] {max-width:100%; overflow:hidden; border-radius:12px;}
        .st-key-individual_dupr_id input {text-transform:uppercase;}
        @media (max-width: 700px) {
            .block-container {padding:.75rem .75rem 2rem;}
            .hero {padding:1.15rem; border-radius:17px;}
            .hero:after {width:120px; height:120px; right:-50px; top:-60px;}
            .hero h1 {padding-right:88px; font-size:1.48rem !important;}
            .hero p {padding-right:0; margin-top:.9rem;}
            .hero-logo-shell {right:.85rem; top:.85rem; transform:none; width:72px; height:65px; padding:3px; border-radius:11px;}
            .hero-logo {border-radius:8px;}
            .rating-grid {grid-template-columns:1fr;}
            div[data-testid="stMetric"] {min-height:90px; padding:.75rem;}
            div[data-testid="stHorizontalBlock"] {flex-wrap:wrap !important; gap:.45rem !important;}
            div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] {
                flex:1 1 calc(50% - .45rem) !important; width:auto !important; min-width:150px !important;
            }
            div[data-testid="stTabs"] button {font-size:.78rem; padding-left:.55rem; padding-right:.55rem; min-height:2.7rem;}
            div[data-testid="stFileUploaderDropzone"] {padding:.65rem !important;}
            div[data-testid="stFileUploaderDropzoneInstructions"] {min-width:0 !important;}
            .stTextInput input, .stSelectbox input {font-size:16px !important;}
            .stButton > button, .stDownloadButton > button {width:100%; min-height:3rem;}
            .section-card, .result-card {padding:.9rem; border-radius:15px;}
            h1 {font-size:1.65rem !important;} h2 {font-size:1.35rem !important;} h3 {font-size:1.08rem !important;}
        }
        @media (max-width: 390px) {
            div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] {
                flex-basis:100% !important; min-width:0 !important;
            }
            .hero p {font-size:.82rem;}
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def get_dupr_token() -> str | None:
    """Load the DUPR token without exposing it in the interface."""
    try:
        token = st.secrets.get("DUPR_TOKEN")
    except (FileNotFoundError, KeyError, TypeError):
        token = None
    token = token or os.environ.get("DUPR_TOKEN")
    return str(token).strip() if token else None


class WindowsTrustStoreAdapter(HTTPAdapter):
    """Use Windows' trusted roots while preserving normal TLS verification."""

    def init_poolmanager(
        self,
        connections: int,
        maxsize: int,
        block: bool = False,
        **pool_kwargs: Any,
    ) -> None:
        context = ssl.create_default_context()
        for store_name in ("ROOT", "CA"):
            for certificate, encoding, _trust in ssl.enum_certificates(store_name):
                if encoding != "x509_asn":
                    continue
                try:
                    context.load_verify_locations(cadata=certificate)
                except ssl.SSLError:
                    # Ignore malformed/unsupported entries and retain all valid roots.
                    continue
        self.poolmanager = PoolManager(
            num_pools=connections,
            maxsize=maxsize,
            block=block,
            ssl_context=context,
            **pool_kwargs,
        )


@st.cache_resource
def get_dupr_session() -> requests.Session:
    """Create a reusable HTTPS session that respects the host OS trust store."""
    session = requests.Session()
    if os.name == "nt" and hasattr(ssl, "enum_certificates"):
        session.mount("https://api.dupr.gg/", WindowsTrustStoreAdapter())
    return session


@st.cache_data(show_spinner=False)
def get_asset_data_uri(relative_path: str) -> str:
    """Return a small local image as an embeddable data URI."""
    asset_path = APP_DIRECTORY / relative_path
    if not asset_path.is_file():
        return ""
    mime_type = "image/jpeg" if asset_path.suffix.lower() in {".jpg", ".jpeg"} else "image/png"
    encoded = base64.b64encode(asset_path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def find_column(df: pd.DataFrame, possible_names: list[str]) -> str | None:
    """Find a Google Forms column despite differences in case and wording."""
    normalized = {str(col).strip().lower(): col for col in df.columns}
    for candidate in possible_names:
        candidate = candidate.strip().lower()
        if candidate in normalized:
            return normalized[candidate]
    for col in df.columns:
        cleaned = str(col).strip().lower()
        if any(
            re.search(rf"\b{re.escape(candidate.strip().lower())}\b", cleaned)
            for candidate in possible_names
        ):
            return col
    return None


def detect_columns(df: pd.DataFrame) -> dict[str, str | None]:
    return {
        "name": find_column(df, ["player name", "full name", "name"]),
        "age": find_column(df, ["age"]),
        "gender": find_column(df, ["gender", "sex"]),
        "phone": find_column(df, ["phone number", "whatsapp", "mobile", "phone"]),
        "category": find_column(df, ["registration category", "categories", "category"]),
        "photo": find_column(df, ["attach one photo", "player photo", "photo"]),
        "payment": find_column(df, ["payment status", "payment"]),
        "dupr": find_column(df, ["dupr id", "dupr"]),
        "skill": find_column(df, ["skill level", "playing level"]),
    }


def clean_dupr_id(value: Any) -> str:
    if pd.isna(value):
        return ""
    cleaned = str(value).strip().upper()
    return cleaned[:-2] if cleaned.endswith(".0") else cleaned


def normalize_gender(value: Any) -> str:
    if pd.isna(value) or not str(value).strip():
        return "Unknown"
    cleaned = str(value).strip().lower().replace(".", "")
    male_values = {"male", "m", "man", "men", "boy", "boys", "boy / man", "boys / men"}
    female_values = {"female", "f", "woman", "women", "girl", "girls", "girl / woman", "girls / women"}
    if cleaned in male_values:
        return "Boys / Men"
    if cleaned in female_values:
        return "Girls / Women"
    return str(value).strip().title()


def parse_age(value: Any) -> int | None:
    try:
        age = float(value)
        if pd.isna(age) or age < 0 or age > 120 or not age.is_integer():
            return None
        return int(age)
    except (TypeError, ValueError):
        return None


def classify_age(age: Any) -> str:
    parsed = parse_age(age)
    if parsed is None:
        return "Unknown"
    if parsed < 18:
        return "U18"
    if parsed <= 30:
        return "18-30"
    return "Above 30"


def safe_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, dict):
        value = value.get("rating") or value.get("value")
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_ratings(player: dict[str, Any] | None) -> tuple[float | None, float | None, float | None]:
    """Return Singles, Doubles, and the higher available classification rating."""
    if not isinstance(player, dict):
        return None, None, None
    ratings = player.get("ratings", {})
    if not isinstance(ratings, dict):
        return None, None, None
    singles = safe_float(ratings.get("singles"))
    doubles = safe_float(ratings.get("doubles"))
    available = [rating for rating in (singles, doubles) if rating is not None]
    return singles, doubles, max(available) if available else None


def classify_level(rating: float | None) -> str:
    if rating is None:
        return "No DUPR Rating"
    return "Advanced" if rating >= ADVANCED_CUTOFF else "Intermediate"


@st.cache_data(ttl=3600, show_spinner=False)
def get_dupr_player(dupr_id: str, token: str) -> dict[str, Any]:
    """Search DUPR by ID using the established public-search request structure."""
    dupr_id = clean_dupr_id(dupr_id)
    if not dupr_id:
        return {"status": "NO_ID", "player": None, "message": "No DUPR ID supplied."}

    body = {
        "offset": 0,
        "limit": 10,
        "query": dupr_id,
        "filter": {"lat": 0, "lng": 0, "radiusInMeters": 50000000},
        "includeUnclaimedPlayers": True,
    }
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
        "Accept": "application/json",
        "User-Agent": "PickleballPlayerManager/1.0",
        "Connection": "close",
    }
    try:
        request_stage = "public search"
        response = get_dupr_session().post(DUPR_API, json=body, headers=headers, timeout=20)
        # DUPR occasionally returns a route-level 404 for public search. Retry
        # once, then use its exact-ID endpoint as a reliable fallback.
        if response.status_code == 404:
            # Close the pooled connection before retrying so a cloud deployment
            # is not pinned to the same unhealthy upstream route.
            get_dupr_session().close()
            get_dupr_session.clear()
            time.sleep(0.4)
            response = get_dupr_session().post(DUPR_API, json=body, headers=headers, timeout=20)
        if response.status_code == 404:
            request_stage = "exact-ID lookup"
            id_response = get_dupr_session().post(
                DUPR_ID_API,
                json={"duprId": dupr_id},
                headers=headers,
                timeout=20,
            )
            if id_response.status_code == 200:
                id_data = id_response.json()
                matches = id_data.get("results", []) if isinstance(id_data, dict) else []
                player_id = matches[0].get("userId") if matches and isinstance(matches[0], dict) else None
                if not player_id:
                    return {"status": "NOT_FOUND", "player": None, "message": "No player matched that DUPR ID."}
                request_stage = "player details"
                response = get_dupr_session().get(
                    DUPR_PLAYER_API.format(player_id=player_id),
                    headers=headers,
                    timeout=20,
                )
                if response.status_code == 200:
                    detail_data = response.json()
                    player = detail_data.get("result") if isinstance(detail_data, dict) else None
                    if isinstance(player, dict) and clean_dupr_id(player.get("duprId", "")) == dupr_id:
                        return {"status": "FOUND", "player": player, "message": "Player found."}
                    return {"status": "BAD_RESPONSE", "player": None, "message": "DUPR returned an unexpected response format."}
            else:
                response = id_response
        if response.status_code == 401:
            return {"status": "TOKEN_EXPIRED", "player": None, "message": "DUPR authentication expired or is invalid."}
        if response.status_code == 403:
            return {"status": "FORBIDDEN", "player": None, "message": "DUPR denied access to this request."}
        if response.status_code != 200:
            return {
                "status": "API_ERROR",
                "player": None,
                "message": f"DUPR {request_stage} returned HTTP {response.status_code}.",
            }
        try:
            data = response.json()
        except ValueError:
            return {"status": "BAD_RESPONSE", "player": None, "message": "DUPR returned an unreadable response."}
        hits = data.get("result", {}).get("hits", []) if isinstance(data, dict) else []
        if not isinstance(hits, list):
            return {"status": "BAD_RESPONSE", "player": None, "message": "DUPR returned an unexpected response format."}
        for player in hits:
            if isinstance(player, dict) and clean_dupr_id(player.get("duprId", "")) == dupr_id:
                return {"status": "FOUND", "player": player, "message": "Player found."}
        return {"status": "NOT_FOUND", "player": None, "message": "No player matched that DUPR ID."}
    except requests.exceptions.Timeout:
        return {"status": "TIMEOUT", "player": None, "message": "The DUPR request timed out."}
    except requests.exceptions.SSLError:
        return {"status": "TLS_ERROR", "player": None, "message": "A secure connection to DUPR could not be established."}
    except requests.exceptions.RequestException:
        return {"status": "NETWORK_ERROR", "player": None, "message": "Could not reach DUPR. Check the network and try again."}
    except Exception:
        return {"status": "API_ERROR", "player": None, "message": "An unexpected DUPR lookup error occurred."}


def player_full_name(player: dict[str, Any] | None) -> str | None:
    if not player:
        return None
    full_name = player.get("fullName")
    if full_name:
        return str(full_name).strip()
    combined = " ".join(str(player.get(part, "")).strip() for part in ("firstName", "lastName")).strip()
    return combined or None


def process_registration(
    row: pd.Series,
    columns: dict[str, str | None],
    token: str,
    dupr_cache: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    dupr_col = columns.get("dupr")
    dupr_id = clean_dupr_id(row.get(dupr_col)) if dupr_col else ""
    if dupr_id:
        if dupr_id not in dupr_cache:
            dupr_cache[dupr_id] = get_dupr_player(dupr_id, token)
            time.sleep(0.08)
        lookup = dupr_cache[dupr_id]
    else:
        lookup = {"status": "NO_ID", "player": None, "message": "No DUPR ID supplied."}

    player = lookup.get("player")
    singles, doubles, classification_rating = extract_ratings(player)
    raw_age = row.get(columns["age"]) if columns.get("age") else None
    raw_gender = row.get(columns["gender"]) if columns.get("gender") else None
    raw_name = row.get(columns["name"]) if columns.get("name") else ""
    raw_phone = row.get(columns["phone"]) if columns.get("phone") else ""
    raw_category = row.get(columns["category"]) if columns.get("category") else ""

    result = row.to_dict()
    result.update(
        {
            "Player Name": "" if pd.isna(raw_name) else str(raw_name).strip(),
            "Entered Age": parse_age(raw_age),
            "Tournament Gender": normalize_gender(raw_gender),
            "Age Division": classify_age(raw_age),
            "DUPR ID": dupr_id,
            "DUPR Verified Name": player_full_name(player),
            "DUPR Singles": singles,
            "DUPR Doubles": doubles,
            "Classification Rating": classification_rating,
            "Tournament Level": classify_level(classification_rating),
            "Registration Category": "" if pd.isna(raw_category) else str(raw_category).strip(),
            "Search Phone": "" if pd.isna(raw_phone) else str(raw_phone).strip(),
            "DUPR Lookup": lookup["status"],
            "DUPR Lookup Message": lookup.get("message", ""),
        }
    )
    return result


def filter_players(
    df: pd.DataFrame,
    age: str = "All",
    level: str = "All",
    gender: str = "All",
    category: str = "All",
    search: str = "",
) -> pd.DataFrame:
    filtered = df.copy()
    if age != "All":
        filtered = filtered[filtered["Age Division"] == age]
    if level != "All":
        filtered = filtered[filtered["Tournament Level"] == level]
    if gender != "All":
        filtered = filtered[filtered["Tournament Gender"] == gender]
    if category != "All":
        filtered = filtered[filtered["Registration Category"].astype(str) == category]
    query = search.strip().lower()
    if query:
        search_blob = (
            filtered["Player Name"].fillna("").astype(str)
            + " " + filtered["DUPR ID"].fillna("").astype(str)
            + " " + filtered["Search Phone"].fillna("").astype(str)
        ).str.lower()
        filtered = filtered[search_blob.str.contains(query, regex=False)]
    return filtered.sort_values("Classification Rating", ascending=False, na_position="last")


def export_order(df: pd.DataFrame) -> pd.DataFrame:
    priority = [
        "Player Name", "Entered Age", "Tournament Gender", "Age Division", "DUPR ID",
        "DUPR Singles", "DUPR Doubles", "Classification Rating", "Tournament Level",
        "Registration Category", "DUPR Verified Name", "DUPR Lookup", "DUPR Lookup Message",
    ]
    columns = [col for col in priority if col in df.columns]
    columns.extend(col for col in df.columns if col not in columns and col != "Search Phone")
    return df[columns]


def write_sheet(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> None:
    export_order(df).to_excel(writer, index=False, sheet_name=sheet_name[:31])


def make_excel(results_df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        write_sheet(writer, results_df, "All Players")
        write_sheet(writer, results_df[results_df["Tournament Level"] == "Advanced"], "All Advanced")
        write_sheet(writer, results_df[results_df["Tournament Level"] == "Intermediate"], "All Intermediate")
        write_sheet(writer, results_df[results_df["Tournament Level"] == "No DUPR Rating"], "Manual Review")

        for age_group in AGE_DIVISIONS:
            for gender in GENDERS:
                for level in ("Advanced", "Intermediate"):
                    division = results_df[
                        (results_df["Age Division"] == age_group)
                        & (results_df["Tournament Gender"] == gender)
                        & (results_df["Tournament Level"] == level)
                    ]
                    if division.empty:
                        continue
                    short_gender = "Boys" if gender == "Boys / Men" else "Girls"
                    write_sheet(writer, division, f"{age_group} {short_gender} {level}")

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="12355B")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 24
            for column_cells in worksheet.columns:
                values = [str(cell.value or "") for cell in column_cells[:200]]
                width = min(max(max(map(len, values), default=10) + 2, 12), 42)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
    output.seek(0)
    return output.getvalue()


def make_filtered_excel(filtered_df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        write_sheet(writer, filtered_df, "Filtered Players")
        worksheet = writer.book["Filtered Players"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="118A74")
            cell.font = Font(color="FFFFFF", bold=True)
        for column_cells in worksheet.columns:
            values = [str(cell.value or "") for cell in column_cells[:200]]
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max(map(len, values), default=10) + 2, 12), 42)
    output.seek(0)
    return output.getvalue()


def metric_count(df: pd.DataFrame, column: str, value: str) -> int:
    return int((df[column] == value).sum()) if column in df.columns else 0


def render_dashboard(results_df: pd.DataFrame | None) -> None:
    st.subheader("Tournament overview")
    if results_df is None or results_df.empty:
        total = advanced = intermediate = manual = boys = girls = 0
        u18 = age_18_30 = above_30 = 0
    else:
        total = len(results_df)
        advanced = metric_count(results_df, "Tournament Level", "Advanced")
        intermediate = metric_count(results_df, "Tournament Level", "Intermediate")
        manual = metric_count(results_df, "Tournament Level", "No DUPR Rating")
        boys = metric_count(results_df, "Tournament Gender", "Boys / Men")
        girls = metric_count(results_df, "Tournament Gender", "Girls / Women")
        u18 = metric_count(results_df, "Age Division", "U18")
        age_18_30 = metric_count(results_df, "Age Division", "18-30")
        above_30 = metric_count(results_df, "Age Division", "Above 30")

    cols = st.columns(3)
    cols[0].metric("Total Players", total)
    cols[1].metric("Boys / Men", boys)
    cols[2].metric("Girls / Women", girls)
    cols = st.columns(3)
    cols[0].metric("Advanced", advanced, f"{advanced / total:.0%}" if total else "0%")
    cols[1].metric("Intermediate", intermediate, f"{intermediate / total:.0%}" if total else "0%")
    found = advanced + intermediate
    cols[2].metric("Manual Review / No Rating", manual, f"DUPR found {found / total:.0%}" if total else "DUPR found 0%")
    cols = st.columns(3)
    cols[0].metric("U18", u18)
    cols[1].metric("18-30", age_18_30)
    cols[2].metric("Above 30", above_30)

    if not total:
        st.info("Upload and process registrations to populate the tournament dashboard.")
        return

    st.markdown("### Quick player list")
    quick_level = st.selectbox(
        "Show Players",
        ["Advanced", "Intermediate", "No DUPR Rating", "All"],
        key="dashboard_quick_level",
    )
    quick_df = results_df if quick_level == "All" else results_df[results_df["Tournament Level"] == quick_level]
    quick_df = quick_df.sort_values("Classification Rating", ascending=False, na_position="last")
    show_player_table(quick_df)

    valid = results_df[
        results_df["Age Division"].isin(AGE_DIVISIONS)
        & results_df["Tournament Gender"].isin(GENDERS)
        & results_df["Tournament Level"].isin(["Advanced", "Intermediate"])
    ]
    st.markdown("### Division summary")
    if valid.empty:
        st.caption("No populated competitive divisions yet.")
    else:
        summary = (
            valid.groupby(["Age Division", "Tournament Gender", "Tournament Level"], observed=True)
            .size().reset_index(name="Players")
        )
        summary["Division"] = summary["Age Division"] + " · " + summary["Tournament Gender"] + " · " + summary["Tournament Level"]
        st.dataframe(summary[["Division", "Players"]], use_container_width=True, hide_index=True)


def show_player_table(df: pd.DataFrame) -> None:
    display_map = {
        "Player Name": "Name", "Entered Age": "Age", "Tournament Gender": "Gender",
        "Age Division": "Age Division", "DUPR ID": "DUPR ID", "DUPR Singles": "DUPR Singles",
        "DUPR Doubles": "DUPR Doubles", "Classification Rating": "Classification Rating",
        "Tournament Level": "Tournament Level", "Registration Category": "Category",
    }
    columns = [col for col in display_map if col in df.columns]
    display = df[columns].rename(columns=display_map)
    st.dataframe(
        display,
        use_container_width=True,
        hide_index=True,
        column_config={
            "DUPR Singles": st.column_config.NumberColumn(format="%.3f"),
            "DUPR Doubles": st.column_config.NumberColumn(format="%.3f"),
            "Classification Rating": st.column_config.NumberColumn(format="%.3f"),
        },
    )


def read_upload(uploaded_file: Any) -> pd.DataFrame:
    if uploaded_file.name.lower().endswith(".csv"):
        try:
            return pd.read_csv(uploaded_file)
        except UnicodeDecodeError:
            uploaded_file.seek(0)
            return pd.read_csv(uploaded_file, encoding="latin-1")
    return pd.read_excel(uploaded_file)


def render_bulk_upload(token: str | None) -> None:
    st.subheader("Bulk registration upload")
    st.caption("Upload the Google Forms export, confirm the detected columns, then process every registration.")
    uploaded_file = st.file_uploader("Registration spreadsheet", type=["xlsx", "csv"], key="registration_upload")
    if uploaded_file is None:
        st.markdown('<div class="section-card"><span class="eyebrow">Accepted formats</span><br><b>Excel (.xlsx)</b> and <b>CSV (.csv)</b></div>', unsafe_allow_html=True)
        return
    try:
        df = read_upload(uploaded_file)
    except Exception as exc:
        st.error(f"Unable to read this spreadsheet. Check that it is a valid XLSX or CSV file. Details: {exc}")
        return
    if df.empty:
        st.error("The spreadsheet has column headers but no registrations.")
        return
    columns = detect_columns(df)
    required_missing = [label for key, label in (("name", "Name"), ("age", "Age"), ("gender", "Gender")) if not columns[key]]
    st.success(f"{len(df):,} registrations detected")
    previous_summary = st.session_state.get("processing_summary")
    if previous_summary:
        st.success(previous_summary["message"])
        if previous_summary["invalid_age"]:
            st.warning(f"{previous_summary['invalid_age']} registration(s) have an invalid age and need correction.")
    detected_count = sum(value is not None for value in columns.values())
    st.caption(f"Detected {detected_count} of {len(columns)} supported fields.")
    with st.expander("Column detection status", expanded=bool(required_missing)):
        labels = {"name": "Name", "age": "Age", "gender": "Gender", "phone": "Phone / WhatsApp", "category": "Category", "photo": "Photo", "payment": "Payment", "dupr": "DUPR ID", "skill": "Skill level"}
        for key, label in labels.items():
            found = columns[key]
            st.markdown(f'<span class="status-pill">{"✓" if found else "—"} {html.escape(label)}: {html.escape(str(found)) if found else "Not detected"}</span>', unsafe_allow_html=True)
    if required_missing:
        st.error("Required columns not detected: " + ", ".join(required_missing) + ". Rename those spreadsheet headers and upload again.")
        return
    if not columns["dupr"]:
        st.warning("No DUPR ID column was detected. You can still process the file; all players will be placed in Manual Review / No DUPR Rating.")
    st.markdown("#### Clean preview")
    preview_cols = [col for col in (columns["name"], columns["age"], columns["gender"], columns["dupr"], columns["category"], columns["phone"]) if col]
    st.dataframe(df[preview_cols].head(10), use_container_width=True, hide_index=True)
    process_clicked = st.button("Process Players", type="primary", use_container_width=True, disabled=not token)
    if not token:
        st.error("Admin configuration required: DUPR_TOKEN is not set. Add it to Streamlit Secrets or the local environment before processing players.")
    if not process_clicked:
        return

    results: list[dict[str, Any]] = []
    dupr_cache: dict[str, dict[str, Any]] = {}
    progress = st.progress(0, text="Preparing registrations…")
    total = len(df)
    for position, (_, row) in enumerate(df.iterrows(), start=1):
        name = str(row.get(columns["name"], "")).strip()
        dupr_id = clean_dupr_id(row.get(columns["dupr"])) if columns["dupr"] else "No ID"
        progress.progress((position - 1) / total, text=f"Checking player {position} of {total} · {name} · {dupr_id or 'No ID'}")
        results.append(process_registration(row, columns, token, dupr_cache))
    progress.progress(1.0, text=f"Processed {total} of {total} players")
    results_df = pd.DataFrame(results)
    st.session_state["results"] = results_df
    st.session_state["source_columns"] = columns

    statuses = results_df["DUPR Lookup"].value_counts()
    found = int(statuses.get("FOUND", 0))
    no_id = int(statuses.get("NO_ID", 0))
    not_found = int(statuses.get("NOT_FOUND", 0))
    api_errors = total - found - no_id - not_found
    invalid_age = int((results_df["Age Division"] == "Unknown").sum())
    st.session_state["processing_summary"] = {
        "message": f"{total} registrations processed · {found} DUPR players found · {no_id} had no DUPR ID · {not_found} not found · {api_errors} API errors",
        "invalid_age": invalid_age,
    }
    st.rerun()


def uppercase_individual_dupr_id() -> None:
    """Keep the visible lookup value in canonical DUPR uppercase form."""
    value = st.session_state.get("individual_dupr_id", "")
    st.session_state["individual_dupr_id"] = str(value).strip().upper()


def render_individual_lookup(token: str | None) -> None:
    st.subheader("Individual DUPR lookup")
    st.caption("Find and classify one player without uploading a spreadsheet.")
    dupr_id = st.text_input(
        "Enter DUPR ID",
        placeholder="N5VJNN",
        key="individual_dupr_id",
        on_change=uppercase_individual_dupr_id,
    ).strip().upper()
    submitted = st.button(
        "Find Player",
        type="primary",
        use_container_width=True,
        disabled=not token,
    )
    if not token:
        st.error("Admin configuration required: DUPR_TOKEN is not set. Ask the app administrator to configure Streamlit Secrets.")
        return
    if not submitted:
        return
    if not dupr_id:
        st.warning("Enter a DUPR ID to search.")
        return
    with st.spinner("Checking DUPR…"):
        lookup = get_dupr_player(dupr_id, token)
    if lookup["status"] != "FOUND":
        friendly = {
            "NOT_FOUND": "No player was found with that DUPR ID. Check the ID and try again.",
            "TOKEN_EXPIRED": "The DUPR token is invalid or expired. Ask the administrator to update DUPR_TOKEN.",
            "FORBIDDEN": "DUPR denied this request. The administrator should check token permissions.",
            "TIMEOUT": "DUPR took too long to respond. Please try again.",
            "NETWORK_ERROR": "DUPR could not be reached. Check the connection and try again.",
            "BAD_RESPONSE": "DUPR returned an unexpected response. Please try again later.",
        }
        st.error(friendly.get(lookup["status"], lookup.get("message", "The lookup could not be completed.")))
        return
    player = lookup["player"]
    singles, doubles, rating = extract_ratings(player)
    level = classify_level(rating)
    badge_class = "advanced" if level == "Advanced" else "intermediate" if level == "Intermediate" else "manual"
    display_level = level.upper() if level != "No DUPR Rating" else "MANUAL REVIEW / NO DUPR RATING"
    fmt = lambda value: f"{value:.3f}" if value is not None else "—"
    st.markdown(
        f"""
        <div class="result-card">
          <div class="eyebrow">DUPR player result</div>
          <div class="result-name">{html.escape(player_full_name(player) or "Player")}</div>
          <div style="color:#627d98;font-weight:650">DUPR ID · {html.escape(clean_dupr_id(player.get("duprId", dupr_id)))}</div>
          <div class="badge {badge_class}">{display_level}</div>
          <div class="rating-grid">
            <div class="rating-box"><div class="rating-label">Singles Rating</div><div class="rating-value">{fmt(singles)}</div></div>
            <div class="rating-box"><div class="rating-label">Doubles Rating</div><div class="rating-value">{fmt(doubles)}</div></div>
            <div class="rating-box"><div class="rating-label">Classification Rating</div><div class="rating-value">{fmt(rating)}</div></div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_players(results_df: pd.DataFrame | None) -> None:
    st.subheader("Players")
    if results_df is None or results_df.empty:
        st.info("Process a registration spreadsheet to search and filter players.")
        return
    search = st.text_input("Search players", placeholder="Name, DUPR ID, or phone number")
    filter_cols = st.columns(2)
    age = filter_cols[0].selectbox("Age Division", ["All"] + AGE_DIVISIONS)
    level = filter_cols[1].selectbox("Level", ["All"] + LEVELS)
    filter_cols = st.columns(2)
    gender = filter_cols[0].selectbox("Gender", ["All"] + GENDERS)
    categories = sorted(value for value in results_df["Registration Category"].dropna().astype(str).unique() if value.strip())
    category = filter_cols[1].selectbox("Registration Category", ["All"] + categories)
    filtered = filter_players(results_df, age, level, gender, category, search)
    st.session_state["filtered_results"] = filtered
    st.markdown(f"### {len(filtered):,} matching player{'s' if len(filtered) != 1 else ''}")
    show_player_table(filtered)
    if not filtered.empty:
        st.caption("Sorted by Classification Rating, highest first.")


def render_export(results_df: pd.DataFrame | None) -> None:
    st.subheader("Export tournament files")
    if results_df is None or results_df.empty:
        st.info("Process registrations before exporting tournament files.")
        return
    filtered = st.session_state.get("filtered_results", results_df)
    st.markdown('<div class="section-card"><span class="eyebrow">Complete workbook</span><br>All players, level lists, manual review, and every populated age / gender / level division.</div>', unsafe_allow_html=True)
    try:
        complete_excel = make_excel(results_df)
        filtered_excel = make_filtered_excel(filtered)
    except Exception as exc:
        st.error(f"Could not create the Excel workbook: {exc}")
        return
    cols = st.columns(2)
    cols[0].download_button(
        "Download Complete Excel", complete_excel, "Pickleball_Tournament_Players.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", use_container_width=True,
    )
    cols[1].download_button(
        f"Download Filtered Excel ({len(filtered)})", filtered_excel, "Pickleball_Filtered_Players.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True,
    )
    st.download_button(
        f"Download Filtered CSV ({len(filtered)})", filtered.to_csv(index=False).encode("utf-8-sig"),
        "Pickleball_Filtered_Players.csv", "text/csv", use_container_width=True,
    )
    with st.expander("Workbook contents"):
        st.markdown("**Always included:** All Players, All Advanced, All Intermediate, Manual Review.\n\n**Included when populated:** U18, 18-30, and Above 30 sheets split by Boys/Girls and Advanced/Intermediate.")


apply_theme()
token = get_dupr_token()
results_df = st.session_state.get("results")
logo_uri = get_asset_data_uri("assets/tournament_logo.jpg")
logo_markup = (
    f'<div class="hero-logo-shell"><img class="hero-logo" src="{logo_uri}" '
    'alt="Vadodara Pickleball League tournament logo"></div>'
    if logo_uri
    else ""
)

st.markdown(
    f"""
    <div class="hero">
      <div class="hero-kicker">Tournament operations</div>
      <h1>🏓 Pickleball Player Manager</h1>
      <p>Registration · DUPR · Tournament Classification</p>
      {logo_markup}
    </div>
    """,
    unsafe_allow_html=True,
)

if token:
    st.caption("● DUPR service configured")
else:
    st.warning("Admin setup needed: DUPR_TOKEN has not been configured. Upload previews remain available, but DUPR lookups and processing are disabled.")

dashboard_tab, upload_tab, lookup_tab, players_tab, export_tab = st.tabs(
    ["Dashboard", "Bulk Upload", "DUPR Lookup", "Players", "Export"]
)
with dashboard_tab:
    render_dashboard(results_df)
with upload_tab:
    render_bulk_upload(token)
with lookup_tab:
    render_individual_lookup(token)
with players_tab:
    render_players(results_df)
with export_tab:
    render_export(results_df)

with st.expander("Classification rules"):
    st.markdown(
        """
        - **Advanced:** classification rating of 4.000 or higher
        - **Intermediate:** classification rating below 4.000
        - **Classification Rating:** the higher available Singles or Doubles DUPR rating
        - **Manual Review / No DUPR Rating:** no ID, no matching player, lookup problem, or player has no rating
        - **Age divisions:** U18 is under 18; 18-30 is inclusive; Above 30 is age 31+
        - **Gender:** Male / M / Boy maps to Boys / Men; Female / F / Girl maps to Girls / Women
        """
    )
